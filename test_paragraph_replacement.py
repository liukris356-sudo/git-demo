import docx

doc = docx.Document()
p_a = doc.add_paragraph('Para A (Keep)')
p_old1 = doc.add_paragraph('Para Old 1 (Delete)')
p_old2 = doc.add_paragraph('Para Old 2 (Delete)')
p_b = doc.add_paragraph('Para B (Keep)')

# Insert new paragraphs before p_b
new_p1 = p_b.insert_paragraph_before('Para New 1')
new_p2 = p_b.insert_paragraph_before('Para New 2')

# Remove old paragraphs
for p in [p_old1, p_old2]:
    p._p.getparent().remove(p._p)

print("Paragraphs in doc after replacement:")
for p in doc.paragraphs:
    print("  -", p.text)
