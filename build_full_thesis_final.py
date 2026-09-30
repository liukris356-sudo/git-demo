# -*- coding: utf-8 -*-
import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn

import thesis_sections_content as tc

def set_run_font(run, ascii_font='Times New Roman', eastasia_font='宋体', size_pt=12, bold=False, italic=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), eastasia_font)

def add_body_before(p_ref, text, indent=True, bold=False):
    p = p_ref.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    run = p.add_run(text)
    set_run_font(run, 'Times New Roman', '宋体', 12, bold=bold)
    return p

def add_heading_before(p_ref, doc, text, level=1):
    p = p_ref.insert_paragraph_before()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(0)
    if level == 1:
        if '章标题' in doc.styles:
            p.style = doc.styles['章标题']
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 16, bold=True)
    elif level == 2:
        if 'Heading 2' in doc.styles:
            p.style = doc.styles['Heading 2']
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 14, bold=True)
    elif level == 3:
        if 'Heading 3' in doc.styles:
            p.style = doc.styles['Heading 3']
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 12.5, bold=True)
    elif level == 4:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '宋体', 12, bold=True)
    return p

def add_equation_before(p_ref, omml_content):
    p = p_ref.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    xml = f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{omml_content}</m:oMathPara>'
    p._p.append(parse_xml(xml))
    return p

def add_image_before(p_ref, doc, img_path, caption, width_in=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return
    p_img = p_ref.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.first_line_indent = Pt(0)
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    
    p_cap = p_ref.insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.first_line_indent = Pt(0)
    run_cap = p_cap.add_run(caption)
    set_run_font(run_cap, 'Times New Roman', '黑体', 10.5, bold=True)

def add_3line_table_before(p_ref, doc, title, headers, data_rows, col_widths=None):
    p_title = p_ref.insert_paragraph_before()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.first_line_indent = Pt(0)
    run_t = p_title.add_run(title)
    set_run_font(run_t, 'Times New Roman', '黑体', 10.5, bold=True)
    
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    hdr_row = table.rows[0]
    for idx, h_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Pt(0)
        run = p.add_run(h_text)
        set_run_font(run, 'Times New Roman', '黑体', 10, bold=True)
        
    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(str(val))
            set_run_font(run, 'Times New Roman', '宋体', 9.5, bold=False)
            
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
    p_ref._p.addprevious(table._tbl)
    
    p_after = p_ref.insert_paragraph_before()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def update_abstract(doc):
    print("Step 1: Updating English Abstract...")
    idx_abs = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == 'ABSTRACT':
            idx_abs = i
            break
            
    if idx_abs is not None:
        p_title = doc.paragraphs[idx_abs]
        p_title.text = tc.ENGLISH_ABSTRACT_TITLE
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(18)
        p_title.paragraph_format.space_after = Pt(14)
        run = p_title.runs[0]
        set_run_font(run, 'Times New Roman', '黑体', 16, bold=True)
        
        paras_to_update = []
        for i in range(idx_abs + 1, min(idx_abs + 15, len(doc.paragraphs))):
            t = doc.paragraphs[i].text.strip()
            if '1.1' in t or '绪论' in t or doc.paragraphs[i].style.name == '章标题':
                break
            paras_to_update.append(doc.paragraphs[i])
            
        all_new_texts = tc.ENGLISH_ABSTRACT_PARAS + [tc.ENGLISH_KEYWORDS]
        for i, new_text in enumerate(all_new_texts):
            if i < len(paras_to_update):
                p = paras_to_update[i]
                p.text = new_text
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Pt(24) if i < len(tc.ENGLISH_ABSTRACT_PARAS) else Pt(0)
                if len(p.runs) > 0:
                    set_run_font(p.runs[0], 'Times New Roman', 'Times New Roman', 12, bold=(i == len(all_new_texts)-1))
            else:
                p_ref = paras_to_update[-1]
                add_body_before(p_ref, new_text, indent=(i < len(tc.ENGLISH_ABSTRACT_PARAS)), bold=(i == len(all_new_texts)-1))
                
        if len(paras_to_update) > len(all_new_texts):
            for p in paras_to_update[len(all_new_texts):]:
                p.text = ''
    print("English Abstract updated successfully.")

def update_chapter1(doc, p_ch2_title):
    print("Step 2: Updating Chapter 1...")
    for i, p in enumerate(doc.paragraphs):
        if '1.4' in p.text and '本章小结' in p.text:
            p_sum = doc.paragraphs[i+1]
            p_sum.text = tc.CH1_SUMMARY
            p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_sum.paragraph_format.line_spacing = 1.5
            p_sum.paragraph_format.first_line_indent = Pt(24)
            if len(p_sum.runs) > 0:
                set_run_font(p_sum.runs[0], 'Times New Roman', '宋体', 12, bold=False)
            break
            
    if p_ch2_title is not None:
        has_1_5 = any('1.5' in p.text and '论文组织结构' in p.text for p in doc.paragraphs)
        if not has_1_5:
            add_heading_before(p_ch2_title, doc, tc.CH1_SECTION5_TITLE, level=2)
            for para_text in tc.CH1_SECTION5_PARAS:
                add_body_before(p_ch2_title, para_text, indent=True)
            p_sep = p_ch2_title.insert_paragraph_before()
            p_sep.paragraph_format.space_before = Pt(14)
    print("Chapter 1 updated successfully.")

def update_chapter2(doc, p_ch2_title):
    print("Step 3: Updating Chapter 2...")
    p_ch2_title.text = '第二章  硬件系统选型与实验台架搭建'
            
    # Heading mappings to ensure strictly logical sequential structure
    ch2_heading_map = {
        '2.2  输电铁塔模拟平台与作业工况分析': '',  # clear orphan heading
        '2.2.1  铁塔模型结构特征与障碍物分布特性': '',  # clear orphan heading
        '2.3  硬件选型': '2.2  硬件系统选型与集成',
        '2.2.1  Jaka机械臂性能参数': '2.2.1  JAKA Zu 7 协作机械臂选型与参数',
        '2.2.2  相机选型': '2.2.2  Intel RealSense D455 深度相机选型与参数',
        '2.2.3  力控传感器选型': '2.2.3  JK-SE-VI-200 六维力/力矩传感器选型与参数',
        '2.3.1 六维力传感器选型': '',  # clear redundant heading
        '表 2-x JK-SE-VI-200 六维力传感器主要技术参数': '表 2-3  JK-SE-VI-200 六维力传感器主要技术参数',
        '表2-4 MPS-M-3000MM-A1性能参数': '',  # remove draw-wire sensor table caption
        '图2-4 MPS-M-3000MM-A1拉绳位移传感器': '',  # remove draw-wire sensor figure caption
        '2.3  输电铁塔模拟平台与作业工况分析': '2.3  输电铁塔模型结构特征与作业工况分析',
        '2.3.1 铁塔模型结构特征与障碍物分布特性': '2.3.1  铁塔模型结构特征与障碍物分布特性',
        '2.3.2  实验台设计': '2.3.2  输电铁塔螺栓作业工况与空间受限约束',
        '2.4  软件平台架构': '2.4  软件控制系统架构设计',
        '2.4.1  ROS操作系统': '2.4.1  ROS 2 软件系统架构与通信机制',
        '2.4.2  ROS2操作系统优势': '2.4.2  多源异构传感器分布式协同与节点设计',
        '2.5  双相机安装分布': '2.5  输电铁塔模拟实验台架与双相机安装设计',
        '2.5.3': '2.5.1  传感器空间拓扑与双相机安装分布'
    }
    
    # Track paragraphs to delete (like the MPS-M sensor image)
    paras_to_remove = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if 'MPS-M-3000MM-A1' in t:
            paras_to_remove.append(p)
            # check if adjacent paragraph has the drawing
            if i > 0 and 'w:drawing' in doc.paragraphs[i-1]._p.xml:
                paras_to_remove.append(doc.paragraphs[i-1])
            if i+1 < len(doc.paragraphs) and 'w:drawing' in doc.paragraphs[i+1]._p.xml:
                paras_to_remove.append(doc.paragraphs[i+1])
        elif t in ch2_heading_map:
            new_val = ch2_heading_map[t]
            p.text = new_val
            if not new_val:
                paras_to_remove.append(p)
            else:
                # set style appropriately
                if new_val.startswith('2.2  ') or new_val.startswith('2.3  ') or new_val.startswith('2.4  ') or new_val.startswith('2.5  '):
                    if 'Heading 2' in doc.styles:
                        p.style = doc.styles['Heading 2']
                elif any(new_val.startswith(f'2.{k}.') for k in range(2, 6)):
                    if 'Heading 3' in doc.styles:
                        p.style = doc.styles['Heading 3']
                        
    for p in set(paras_to_remove):
        try:
            p._p.getparent().remove(p._p)
        except Exception:
            pass
            
    # Clean Table 3 duplicate text
    if len(doc.tables) > 3:
        t3 = doc.tables[3]
        clean_map = {
            '±8 N⋅m±8 N⋅m': '±8 N·m',
            '0.5% F.S.0.5% F.S.': '0.5% F.S.',
            '≥300% F.S.≥300% F.S.': '≥300% F.S.',
            'DC 12 VDC 12 V': 'DC 12 V',
            '5 ∘C∼80 ∘C5 ∘C∼80 ∘C': '5 ℃ ~ 80 ℃',
            '力矩量程（Mx/My/MzMx\u200b/My\u200b/Mz\u200b）': '力矩量程（Mx/My/Mz）'
        }
        for row in t3.rows:
            for cell in row.cells:
                for k, v in clean_map.items():
                    if k in cell.text:
                        cell.text = cell.text.replace(k, v)
    print("Chapter 2 updated successfully.")

def update_chapter3(doc, p_ch3_title):
    print("Step 4: Updating Chapter 3...")
    p_ch3_title.text = '第三章  障碍物与螺栓位姿双视觉识别方法'
            
    ch3_heading_map = {
        '3.1  单视觉的场景瓶颈分析': '3.2  单视觉的场景瓶颈分析',
        '3.1.1  桁架遮挡导致的目标漏检问题': '3.2.1  桁架遮挡导致的目标漏检问题',
        '3.1.2  全局视野与局部精度的矛盾': '3.2.2  全局视野与局部精度的矛盾',
        '3.2  双视觉协同感知架构设计': '3.3  双视觉协同感知架构设计',
        '3.2.1  全局障碍感知模块': '3.3.1  全局障碍感知模块',
        '3.2.2  局部螺栓精定位模块': '3.3.2  局部螺栓精定位模块',
        '3.3  障碍物检测与空间映射方法': '3.4  障碍物检测与空间映射方法',
        '3.4  螺栓位姿识别与坐标解算': '3.5  螺栓位姿识别与坐标解算',
        '3.5  识别定位实验与分析': '3.6  识别定位实验与分析',
        '3.5.1  障碍物检测准确率': '3.6.1  障碍物检测准确率',
        '3.5.2  螺栓位姿识别精度': '3.6.2  螺栓位姿识别精度',
        '3.6  本章小结': '3.7  本章小结'
    }
    
    for p in doc.paragraphs:
        t = p.text.strip()
        for k, v in ch3_heading_map.items():
            if t == k:
                p.text = v
                if v.startswith('3.2  ') or v.startswith('3.3  ') or v.startswith('3.4  ') or v.startswith('3.5  ') or v.startswith('3.6  ') or v.startswith('3.7  '):
                    if 'Heading 2' in doc.styles:
                        p.style = doc.styles['Heading 2']
                elif any(v.startswith(f'3.{k}.') for k in range(2, 7)):
                    if 'Heading 3' in doc.styles:
                        p.style = doc.styles['Heading 3']
                break
                
    # Embed Images in Chapter 3
    img3_1 = '/home/liu/projects/git-demo-admittance/fig3_1_bottlenecks.png'
    img3_2 = '/home/liu/projects/git-demo-admittance/fig3_2_architecture.png'
    img3_3 = '/home/liu/projects/git-demo-admittance/fig3_3_pose_pipeline.png'
    
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if '图 3-1 铁塔复杂场景下单视觉感知瓶颈机理分析' in t:
            prev_p = doc.paragraphs[i-1]
            if 'w:drawing' not in prev_p._p.xml and os.path.exists(img3_1):
                p_img = p.insert_paragraph_before()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.first_line_indent = Pt(0)
                run = p_img.add_run()
                run.add_picture(img3_1, width=Inches(5.8))
        elif '图 3-2 双视觉协同感知系统总体架构与数据流向图' in t:
            prev_p = doc.paragraphs[i-1]
            if 'w:drawing' not in prev_p._p.xml and os.path.exists(img3_2):
                p_img = p.insert_paragraph_before()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.first_line_indent = Pt(0)
                run = p_img.add_run()
                run.add_picture(img3_2, width=Inches(5.8))
        elif '图 3-3 局部视觉螺栓孔位姿解算与几何拟合原理流程图' in t:
            prev_p = doc.paragraphs[i-1]
            if 'w:drawing' not in prev_p._p.xml and os.path.exists(img3_3):
                p_img = p.insert_paragraph_before()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.first_line_indent = Pt(0)
                run = p_img.add_run()
                run.add_picture(img3_3, width=Inches(5.8))
                
    # Style Table 4 and Table 5 as 3-line tables
    if len(doc.tables) > 5:
        for t_idx in [4, 5]:
            t = doc.tables[t_idx]
            tblPr = t._tbl.tblPr
            tblBorders = parse_xml(r'''
                <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
                    <w:left w:val="none"/>
                    <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
                    <w:right w:val="none"/>
                    <w:insideH w:val="none"/>
                    <w:insideV w:val="none"/>
                </w:tblBorders>
            ''')
            tblPr.append(tblBorders)
            for cell in t.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(r'''
                    <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                    </w:tcBorders>
                ''')
                tcPr.append(tcBorders)
    print("Chapter 3 updated successfully.")

def update_chapter4(doc, p_ch4_title, p_ch5_title):
    print("Step 5: Updating Chapter 4...")
    p_ch4_title.text = '第四章  输电铁塔狭窄空间避障路径规划算法'
            
    caption_map = {
        '图4-1 P3P投影示意图': '图 4-1  机械臂与障碍物碰撞检测包围盒模型示意图',
        '图2  APF-Informed-RRT*算法流程图': '图 4-2  APF-Informed-RRT*算法流程图',
        '图3  混合采样（高斯采样+目标偏置采样）示意图': '图 4-3  混合采样（高斯采样+目标偏置采样）示意图',
        '图4  椭圆域采样过程': '图 4-4  椭圆域采样过程',
        '图5  靠近目标点时斥力减小': '图 4-5  靠近目标点时斥力减小',
        '图6  安全环境下动态权重效果': '图 4-6  安全环境下动态权重效果',
        '图7  危险环境下动态权重效果': '图 4-7  危险环境下动态权重效果',
        '图8  二维单目标点算法对比试验': '图 4-8  二维单目标点算法对比试验',
        '表1\u3000二维单目标点算性能对比': '表 4-1  二维单目标点算法性能对比',
        '图9  三维单目标点算法对比试验': '图 4-9  三维单目标点算法对比试验',
        '表2\u3000三维单目标点算性能对比': '表 4-2  三维复杂空间单目标点算法性能对比',
        '图10  三维复杂空间多目标点路径规划对比实验': '图 4-10  三维复杂空间多目标点路径规划对比实验',
        '图11  三种算法效果对比': '图 4-11  三种算法规划效果与节点分布对比',
        '表3\u3000三种算法性能对比': '表 4-3  三维复杂空间多目标点算法性能对比',
        '图12  参数敏感度分析': '图 4-12  目标偏置概率对规划性能的影响',
        '图13  参数敏感度分析': '图 4-13  椭圆域收缩率对规划性能的影响',
        '图14  不同算法收敛性能对比': '图 4-14  不同算法在密集障碍环境下的收敛性能对比',
        '图15 铁塔作业区域及位置': '图 4-15  输电铁塔作业区域与目标螺栓点位分布',
        '图16 物理仿真环境设置': '图 4-16  PyBullet 物理仿真环境搭建与坐标系标定',
        '表4\u3000铁塔仿真环境坐标设置': '表 4-4  输电铁塔物理仿真环境空间坐标设置',
        '图17 机械臂避障路径规划全流程': '图 4-17  机械臂多螺栓连续避障规划运动包络全流程',
        '表5\u3000算法性能数据': '表 4-5  输电铁塔物理仿真多阶段规划算法性能数据',
        '2.1.1 目标偏置采样': '4.3.1.1  目标偏置采样',
        '2.1.2基于高斯扰动的视觉引导采样': '4.3.1.2  基于高斯扰动的视觉引导采样',
        '4.3.2  基于Informed椭圆域的收缩采样': '4.3.1.3  基于Informed椭圆域的收缩采样',
        '2.2.1 引力势场': '4.3.2.1  引力势场设计',
        '2.2.2斥力势场及改进': '4.3.2.2  改进GNRON斥力势场',
        '2.2.3基于障碍物距离的动态权重调节': '4.3.2.3  基于障碍物距离的动态权重调节',
        '2.3.1基于贪婪策略的路径剪枝': '4.3.3.1  基于贪婪策略的路径剪枝',
        '2.3.2基于三次B样条的轨迹平滑': '4.3.3.2  基于三次B样条的轨迹平滑',
        '实验验证和结果分析': '4.4  算法仿真与对比分析',
        '二维复杂环境单目标点路径规划': '4.4.1  二维复杂环境单目标点路径规划仿真',
        '三维复杂空间单目标点路径规划': '4.4.2  三维复杂空间单目标点路径规划仿真',
        '三维复杂空间多目标点路径规划': '4.4.3  三维复杂空间多目标点路径规划仿真',
        '参数敏感性分析': '4.4.4  参数敏感性分析',
        '算法收敛性能分析': '4.4.5  算法收敛性能分析',
        '输电铁塔平台验证与分析': '4.4.6  输电铁塔物理仿真验证与分析',
        '3.6.1实验场景与任务设定': '4.4.6.1  铁塔物理仿真场景与任务设定',
        '3.6.2 仿真数据统计与性能分析': '4.4.6.2  仿真数据统计与性能分析'
    }
    
    for p in doc.paragraphs:
        t = p.text.strip()
        for k, v in caption_map.items():
            if k == t:
                p.text = v
                if '4.4.' in v and len(v) < 30:
                    if 'Heading 2' in doc.styles:
                        p.style = doc.styles['Heading 2']
                elif '4.3.' in v and len(v) < 30:
                    if 'Heading 3' in doc.styles:
                        p.style = doc.styles['Heading 3']
                break
                
    if p_ch5_title is not None:
        has_ch4_sum = any('4.5' in p.text and '本章小结' in p.text for p in doc.paragraphs)
        if not has_ch4_sum:
            add_heading_before(p_ch5_title, doc, tc.CH4_SUMMARY_TITLE, level=2)
            for para_text in tc.CH4_SUMMARY_PARAS:
                add_body_before(p_ch5_title, para_text, indent=True)
            p_sep = p_ch5_title.insert_paragraph_before()
            p_sep.paragraph_format.space_before = Pt(14)
    print("Chapter 4 updated successfully.")

def update_chapter5(doc, p_ch5_title):
    print("Step 6: Updating Chapter 5...")
    p_ch5_title.text = '第五章  螺栓柔顺装配力控策略'
            
    ch5_heading_map = {
        '5.2  铁塔螺栓装配的力控需求分析': '5.2  铁塔螺栓装配的力控需求分析',
        '5.1.1  螺栓装配的接触力约束': '5.2.1  螺栓装配的接触力学约束与六阶段力学演变',
        '5.1.2  传统位置控制的局限性': '5.2.2  传统位置控制的局限性',
        '5.2  基于导纳控制的柔顺装配策略': '5.3  基于导纳控制的柔顺装配策略',
        '5.2.1  末端力信息处理与重力补偿': '5.3.1  末端力信息处理与重力补偿',
        '5.2.2  位置-力混合控制律设计': '5.3.2  位置-力混合控制律设计',
        '5.3  视觉引导的力控装配流程': '5.4  视觉引导的力控装配流程',
        '5.3.1  视觉预定位阶段': '5.4.1  视觉预定位阶段',
        '5.3.2  力控柔顺插入阶段': '5.4.2  力控柔顺插入阶段',
        '5.4  本章小结': '5.5  本章小结'
    }
    
    for p in doc.paragraphs:
        t = p.text.strip()
        for k, v in ch5_heading_map.items():
            if t == k:
                p.text = v
                if v.startswith('5.2  ') or v.startswith('5.3  ') or v.startswith('5.4  ') or v.startswith('5.5  '):
                    if 'Heading 2' in doc.styles:
                        p.style = doc.styles['Heading 2']
                elif any(v.startswith(f'5.{k}.') for k in range(2, 5)):
                    if 'Heading 3' in doc.styles:
                        p.style = doc.styles['Heading 3']
                break
                
    # Remove stray sentence
    for p in doc.paragraphs:
        if '式中R表示螺栓目标检测识别率' in p.text:
            p.text = ''
            
    # Restore formulas in 6-stage contact mechanics
    stage_a = '（a）初始表面点接触阶段（状态 P1）\n在全局与局部视觉粗定位引导下，机械臂将螺栓运送至基体件（铁塔节点板或角钢）表面。由于视觉标定与测量残差的存在，螺栓端部并未直接对准孔心，而是以微小的倾斜姿态与基体上表面发生单点接触。此时螺栓受到重力 G、基体表面法向支撑力 N1、表面摩擦力 f1 以及机械臂末端的主动驱动力 FA1 的共同作用；机械臂需根据接触产生的偏心反力矩施加自适应顺应力矩 Mc1，使螺栓向孔口方向微调并滑动。'
    stage_b = '（b）平面滑移对心阶段（状态 P2）\n螺栓在机械臂末端下压与横向驱动力 FA2 作用下，克服滑动摩擦阻力 f2 沿基体表面向预制孔中心匀速滑动。该阶段垂直支撑力 N2 与轴向下压力维持动态平衡，力控系统需严格抑制轴向力在安全接触阈值内（10 ~ 20 N），既保证螺栓不脱离基体表面，又避免压力过大刮伤工件表面涂层。'
    stage_c = '（c）孔口落入与倒角两点接触阶段（状态 P3）\n螺栓下端滑至孔口倒角边缘处受阻力突变开始“掉入”孔口。由于轴线尚未与孔轴线重合，螺栓头部与孔壁两侧发生倾斜接触，形成初始的两点接触约束。此时，孔壁两侧对螺栓产生法向反力 N3a、N3b 及切向摩擦阻力 f3a、f3b，若下压力 FA3 过大将极易在孔口边缘诱发“几何楔紧（Wedging）”。因此，机械臂必须依靠六维力传感器感知到的力矩突变，启动主动自适应调姿力矩 Mc3，驱动螺栓绕接触点旋转“摆正”，消除轴线夹角。'
    stage_d = '（d）深孔两点接触与自适应纠偏阶段（状态 P4）\n螺栓部分杆身深入预制孔内，此时螺栓下端外壁与孔口对侧内壁同时紧贴，处于典型的深孔两点接触状态。该阶段是装配过程中最危险的“卡滞（Jamming）”高发期：孔壁两侧产生的法向力 N4a、N4b 和纵向摩擦力 f4a、f4b 构成了强烈的阻力偶矩，阻碍螺栓继续下潜。此时常规位置控制若强行推进必将导致螺栓咬死；系统必须通过导纳控制算法，利用测得的翻转弯矩实时微调机械臂末端横向位置与空间偏角，迫使螺栓与孔壁法线方向严格重合。'
    stage_e = '（e）同轴垂直平稳插入阶段（状态 P5）\n随着轴线偏角的彻底消除（θ -> 0），螺栓与预制孔恢复同轴共线状态，原有的两点刚性卡滞退化为均匀的柱面滑动摩擦接触。此时横向干涉力与翻转力矩均衰减至零附近，机械臂只需施加恒定的轴向推进推力 FA5 克服平稳滑动阻力 f5，驱动螺栓快速、平稳地向孔底深处匀速滑入。'
    stage_f = '（f）底部贴合与装配完成阶段（状态 P6）\n当螺栓头部法兰面与基体件表面完全密合时，由于基体件孔底具有极高的法向刚度，垂直支撑反力 N6 发生阶跃式剧增。控制器实时捕获到该轴向力突变并判断其超过设定的装配终止阈值（Fstop = 45 N），立即切断轴向下压伺服指令并锁定当前位姿，宣告全流程柔顺装配顺利完成。'
    
    stage_idx = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('（a）初始表面点接触阶段') and stage_idx == 0:
            p.text = stage_a
            stage_idx = 1
        elif t.startswith('（b）平面滑移对心阶段') and stage_idx == 1:
            p.text = stage_b
            stage_idx = 2
        elif t.startswith('（c）孔口落入与倒角两点接触阶段') and stage_idx == 2:
            p.text = stage_c
            stage_idx = 3
        elif t.startswith('（d）深孔两点接触与自适应纠偏阶段') and stage_idx == 3:
            p.text = stage_d
            stage_idx = 4
        elif t.startswith('（e）同轴垂直平稳插入阶段') and stage_idx == 4:
            p.text = stage_e
            stage_idx = 5
        elif t.startswith('（f）底部贴合与装配完成阶段') and stage_idx == 5:
            p.text = stage_f
            stage_idx = 6
            
    # Embed Fig 5-1
    img5_1 = '/home/liu/.gemini/antigravity/brain/459352ac-3327-40c2-bb29-4319774ed722/.user_uploaded/media_1790058152907.png'
    for i, p in enumerate(doc.paragraphs):
        if '如图 5-x 所示' in p.text or '六个典型的力学演化阶段' in p.text:
            p.text = p.text.replace('如图 5-x 所示', '如图 5-1 所示')
            next_p = doc.paragraphs[i+1]
            if 'w:drawing' not in next_p._p.xml and os.path.exists(img5_1):
                p_img = next_p.insert_paragraph_before()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.first_line_indent = Pt(0)
                run = p_img.add_run()
                run.add_picture(img5_1, width=Inches(5.8))
                
                p_cap = next_p.insert_paragraph_before()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(8)
                p_cap.paragraph_format.first_line_indent = Pt(0)
                run_cap = p_cap.add_run('图 5-1  螺栓柔顺装配轴孔配合力学状态演变与受力分析图')
                set_run_font(run_cap, 'Times New Roman', '黑体', 10.5, bold=True)
            break
    print("Chapter 5 updated successfully.")

def update_chapter6_and_7(doc):
    print("Step 7: Creating Chapter 6 and Chapter 7...")
    p_old_ch6 = None
    for p in doc.paragraphs:
        if '总结与展望' in p.text and p.style.name == '章标题':
            p_old_ch6 = p
            break
            
    p_ref_start = None
    for p in doc.paragraphs:
        if '参考文献' in p.text and p.style.name == '章标题':
            p_ref_start = p
            break
            
    if p_old_ch6 is not None and p_ref_start is not None:
        # Collect paragraphs using XML identity
        paras_to_remove = []
        collect = False
        for p in doc.paragraphs:
            if p._p is p_old_ch6._p:
                collect = True
            if p._p is p_ref_start._p:
                break
            if collect:
                paras_to_remove.append(p)
                
        # Insert Chapter 6 before p_ref_start
        add_heading_before(p_ref_start, doc, tc.CH6_TITLE, level=1)
        add_body_before(p_ref_start, tc.CH6_INTRO, indent=True)
        
        # 6.1
        add_heading_before(p_ref_start, doc, tc.CH6_1_TITLE, level=2)
        for para in tc.CH6_1_PARAS:
            add_body_before(p_ref_start, para, indent=True)
            
        # 6.2
        add_heading_before(p_ref_start, doc, tc.CH6_2_TITLE, level=2)
        add_heading_before(p_ref_start, doc, tc.CH6_2_1_TITLE, level=3)
        for para in tc.CH6_2_1_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_1_TITLE, tc.TABLE_6_1_HEADERS, tc.TABLE_6_1_DATA)
        
        add_heading_before(p_ref_start, doc, tc.CH6_2_2_TITLE, level=3)
        for para in tc.CH6_2_2_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_2_TITLE, tc.TABLE_6_2_HEADERS, tc.TABLE_6_2_DATA)
        
        # 6.3
        add_heading_before(p_ref_start, doc, tc.CH6_3_TITLE, level=2)
        for para in tc.CH6_3_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_3_TITLE, tc.TABLE_6_3_HEADERS, tc.TABLE_6_3_DATA)
        
        # 6.4
        add_heading_before(p_ref_start, doc, tc.CH6_4_TITLE, level=2)
        add_heading_before(p_ref_start, doc, tc.CH6_4_1_TITLE, level=3)
        for para in tc.CH6_4_1_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_4_TITLE, tc.TABLE_6_4_HEADERS, tc.TABLE_6_4_DATA)
        
        add_heading_before(p_ref_start, doc, tc.CH6_4_2_TITLE, level=3)
        for para in tc.CH6_4_2_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_5_TITLE, tc.TABLE_6_5_HEADERS, tc.TABLE_6_5_DATA)
        
        # 6.5
        add_heading_before(p_ref_start, doc, tc.CH6_5_TITLE, level=2)
        for para in tc.CH6_5_PARAS:
            add_body_before(p_ref_start, para, indent=True)
        add_3line_table_before(p_ref_start, doc, tc.TABLE_6_6_TITLE, tc.TABLE_6_6_HEADERS, tc.TABLE_6_6_DATA)
        
        # 6.6
        add_heading_before(p_ref_start, doc, tc.CH6_6_TITLE, level=2)
        for para in tc.CH6_6_PARAS:
            add_body_before(p_ref_start, para, indent=True)
            
        p_sep1 = p_ref_start.insert_paragraph_before()
        p_sep1.paragraph_format.space_before = Pt(14)
        
        # Chapter 7
        add_heading_before(p_ref_start, doc, tc.CH7_TITLE, level=1)
        add_heading_before(p_ref_start, doc, tc.CH7_1_TITLE, level=2)
        for para in tc.CH7_1_PARAS:
            add_body_before(p_ref_start, para, indent=True)
            
        add_heading_before(p_ref_start, doc, tc.CH7_2_TITLE, level=2)
        for para in tc.CH7_2_PARAS:
            add_body_before(p_ref_start, para, indent=True)
            
        p_sep2 = p_ref_start.insert_paragraph_before()
        p_sep2.paragraph_format.space_before = Pt(14)
        
        # Delete old paragraphs
        for p in paras_to_remove:
            try:
                p._p.getparent().remove(p._p)
            except Exception:
                pass
                
        print(f"Removed {len(paras_to_remove)} old paragraphs and inserted new Chapter 6 & 7.")

def main():
    base_file = '/home/liu/文档/基于视觉感知的机械臂避障路径规划---刘庆-----毕业论文(1)(1).docx'
    output_doc_path = '/home/liu/文档/基于视觉感知的机械臂避障路径规划---刘庆-----毕业论文_修订完成版.docx'
    output_dl_path = '/home/liu/下载/基于视觉感知的机械臂避障路径规划---刘庆-----毕业论文_修订完成版.docx'
    
    print(f"Loading document: {base_file}...")
    doc = Document(base_file)
    
    def get_title_para(intro_prefix):
        idx = None
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip().startswith(intro_prefix):
                idx = i
                break
        if idx is not None:
            for i in range(idx - 1, -1, -1):
                if doc.paragraphs[i].style.name == '章标题':
                    return doc.paragraphs[i]
        return None

    p_ch1_title = get_title_para('1.1')
    p_ch2_title = get_title_para('2.1')
    p_ch3_title = get_title_para('3.1')
    p_ch4_title = get_title_para('4.1')
    p_ch5_title = get_title_para('5.1')

    if p_ch1_title is not None:
        p_ch1_title.text = '第一章  绪论'
        
    update_abstract(doc)
    update_chapter1(doc, p_ch2_title)
    update_chapter2(doc, p_ch2_title)
    update_chapter3(doc, p_ch3_title)
    update_chapter4(doc, p_ch4_title, p_ch5_title)
    update_chapter5(doc, p_ch5_title)
    update_chapter6_and_7(doc)
    
    print(f"Saving to {output_doc_path}...")
    doc.save(output_doc_path)
    
    print(f"Copying to {output_dl_path}...")
    shutil.copy2(output_doc_path, output_dl_path)
    
    print("SUCCESS! All chapters revised and exported.")

if __name__ == '__main__':
    main()
