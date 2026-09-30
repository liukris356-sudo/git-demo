import os
import sys
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

import thesis_sections_content as tc

def set_run_font(run, ascii_font='Times New Roman', eastasia_font='宋体', size_pt=12, bold=False, italic=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), eastasia_font)

def add_body_before(p_ref, text, indent=True, bold=False):
    p = p_ref.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    run = p.add_run(text)
    set_run_font(run, 'Times New Roman', '宋体', 12, bold=bold)
    return p

def add_heading_before(p_ref, text, level=1):
    p = p_ref.insert_paragraph_before()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(0)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 16, bold=True)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 14, bold=True)
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '黑体', 12.5, bold=True)
    elif level == 4:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        set_run_font(run, 'Times New Roman', '宋体', 12, bold=True)
    return p

def add_equation_before(p_ref, omml_content):
    p = p_ref.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    xml = f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{omml_content}</m:oMathPara>'
    p._p.append(parse_xml(xml))
    return p

def add_image_before(p_ref, doc, img_path, caption, width_in=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return
    p_img = p_ref.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.first_line_indent = Pt(0)
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    
    p_cap = p_ref.insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.first_line_indent = Pt(0)
    run_cap = p_cap.add_run(caption)
    set_run_font(run_cap, 'Times New Roman', '黑体', 10.5, bold=True)

def add_3line_table_before(p_ref, doc, title, headers, data_rows, col_widths=None):
    p_title = p_ref.insert_paragraph_before()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.first_line_indent = Pt(0)
    run_t = p_title.add_run(title)
    set_run_font(run_t, 'Times New Roman', '黑体', 10.5, bold=True)
    
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Apply standard 3-line borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    # Format header row
    hdr_row = table.rows[0]
    for idx, h_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Pt(0)
        run = p.add_run(h_text)
        set_run_font(run, 'Times New Roman', '黑体', 10, bold=True)
        
    # Format data rows
    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(str(val))
            set_run_font(run, 'Times New Roman', '宋体', 9.5, bold=False)
            
    # Set column widths if provided
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
    p_ref._p.addprevious(table._tbl)
    
    # Add small spacing after table
    p_after = p_ref.insert_paragraph_before()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

print("Helpers ready.")
