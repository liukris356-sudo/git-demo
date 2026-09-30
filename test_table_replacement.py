import docx

doc = docx.Document()
p_a = doc.add_paragraph('Para A (Keep)')
t_old = doc.add_table(rows=1, cols=1)
t_old.cell(0, 0).text = 'Old Table'
p_old = doc.add_paragraph('Old text')
p_b = doc.add_paragraph('Para B (Keep)')

# Insert new table before p_b
t_new = doc.add_table(rows=1, cols=1)
t_new.cell(0, 0).text = 'New Table'
p_b._p.addprevious(t_new._tbl)

# Remove old table and old para
t_old._tbl.getparent().remove(t_old._tbl)
p_old._p.getparent().remove(p_old._p)

print("Result paragraphs:")
for p in doc.paragraphs:
    print("  p:", p.text)
print("Result tables:")
for t in doc.tables:
    print("  t:", t.cell(0, 0).text)
